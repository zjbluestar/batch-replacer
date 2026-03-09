# -*- coding: utf-8 -*-
"""
批量关键词替换工具 v1.1
功能：批量修改Word/Excel文件中的关键词
支持格式：.docx, .xlsx
说明：不支持旧格式 .doc/.xls，请先用Office另存为新格式
"""

import os
import sys
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from pathlib import Path

# 尝试导入文档处理库
try:
    from docx import Document
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    from openpyxl import load_workbook
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


class BatchReplacerApp:
    def __init__(self, root):
        self.root = root
        self.root.title("批量关键词替换工具 v1.1")
        self.root.geometry("600x500")
        self.root.resizable(True, True)
        
        # 替换规则列表
        self.replace_rules = []
        
        self.create_widgets()
        self.check_dependencies()
    
    def check_dependencies(self):
        """检查依赖库"""
        missing = []
        if not HAS_DOCX:
            missing.append("python-docx (处理.docx)")
        if not HAS_OPENPYXL:
            missing.append("openpyxl (处理.xlsx)")
        
        if missing:
            self.log(f"⚠️ 缺少依赖库: {', '.join(missing)}")
        else:
            self.log("✓ 依赖库检查通过")
            self.log("支持格式: .docx, .xlsx")
            self.log("注意: .doc/.xls 请先用Office另存为新格式")
    
    def create_widgets(self):
        """创建界面组件"""
        # 主框架
        main_frame = ttk.Frame(self.root, padding="10")
        main_frame.pack(fill=tk.BOTH, expand=True)
        
        # 文件夹选择
        folder_frame = ttk.LabelFrame(main_frame, text="选择文件夹", padding="5")
        folder_frame.pack(fill=tk.X, pady=5)
        
        self.folder_path = tk.StringVar()
        ttk.Entry(folder_frame, textvariable=self.folder_path, width=50).pack(side=tk.LEFT, padx=5)
        ttk.Button(folder_frame, text="浏览...", command=self.select_folder).pack(side=tk.LEFT)
        
        # 替换规则
        rules_frame = ttk.LabelFrame(main_frame, text="替换规则", padding="5")
        rules_frame.pack(fill=tk.BOTH, expand=True, pady=5)
        
        # 规则输入
        input_frame = ttk.Frame(rules_frame)
        input_frame.pack(fill=tk.X, pady=5)
        
        ttk.Label(input_frame, text="旧词:").pack(side=tk.LEFT, padx=5)
        self.old_word = tk.StringVar()
        ttk.Entry(input_frame, textvariable=self.old_word, width=20).pack(side=tk.LEFT, padx=5)
        
        ttk.Label(input_frame, text="→").pack(side=tk.LEFT)
        
        ttk.Label(input_frame, text="新词:").pack(side=tk.LEFT, padx=5)
        self.new_word = tk.StringVar()
        ttk.Entry(input_frame, textvariable=self.new_word, width=20).pack(side=tk.LEFT, padx=5)
        
        ttk.Button(input_frame, text="添加规则", command=self.add_rule).pack(side=tk.LEFT, padx=10)
        
        # 规则列表
        list_frame = ttk.Frame(rules_frame)
        list_frame.pack(fill=tk.BOTH, expand=True, pady=5)
        
        columns = ("old", "new")
        self.rules_tree = ttk.Treeview(list_frame, columns=columns, show="headings", height=5)
        self.rules_tree.heading("old", text="旧词")
        self.rules_tree.heading("new", text="新词")
        self.rules_tree.column("old", width=200)
        self.rules_tree.column("new", width=200)
        
        scrollbar = ttk.Scrollbar(list_frame, orient=tk.VERTICAL, command=self.rules_tree.yview)
        self.rules_tree.configure(yscrollcommand=scrollbar.set)
        
        self.rules_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        # 删除规则按钮
        btn_frame = ttk.Frame(rules_frame)
        btn_frame.pack(fill=tk.X, pady=5)
        ttk.Button(btn_frame, text="删除选中规则", command=self.delete_rule).pack(side=tk.RIGHT)
        
        # 输出选项
        output_frame = ttk.LabelFrame(main_frame, text="输出选项", padding="5")
        output_frame.pack(fill=tk.X, pady=5)
        
        self.output_option = tk.StringVar(value="new_folder")
        ttk.Radiobutton(output_frame, text="输出到新文件夹", variable=self.output_option, value="new_folder").pack(side=tk.LEFT, padx=10)
        ttk.Radiobutton(output_frame, text="直接覆盖原文件", variable=self.output_option, value="overwrite").pack(side=tk.LEFT, padx=10)
        
        # 日志区域
        log_frame = ttk.LabelFrame(main_frame, text="处理日志", padding="5")
        log_frame.pack(fill=tk.BOTH, expand=True, pady=5)
        
        self.log_text = tk.Text(log_frame, height=8, state=tk.DISABLED)
        log_scrollbar = ttk.Scrollbar(log_frame, orient=tk.VERTICAL, command=self.log_text.yview)
        self.log_text.configure(yscrollcommand=log_scrollbar.set)
        
        self.log_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        log_scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        # 操作按钮（底部醒目位置）
        button_frame = ttk.Frame(main_frame)
        button_frame.pack(fill=tk.X, pady=15)
        
        # 开始替换按钮（大而醒目）
        self.start_btn = tk.Button(
            button_frame, 
            text="▶ 开始执行替换", 
            command=self.start_replace,
            font=("微软雅黑", 14, "bold"),
            bg="#4CAF50",
            fg="white",
            width=20,
            height=2,
            relief=tk.RAISED,
            cursor="hand2"
        )
        self.start_btn.pack(side=tk.LEFT, padx=20)
        
        ttk.Button(button_frame, text="退出程序", command=self.root.quit).pack(side=tk.RIGHT, padx=20)
    
    def select_folder(self):
        """选择文件夹"""
        folder = filedialog.askdirectory(title="选择包含Word/Excel文件的文件夹")
        if folder:
            self.folder_path.set(folder)
            self.log(f"已选择文件夹: {folder}")
    
    def add_rule(self):
        """添加替换规则"""
        old = self.old_word.get().strip()
        new = self.new_word.get().strip()
        
        if not old:
            messagebox.showwarning("警告", "请输入旧词")
            return
        
        self.replace_rules.append((old, new))
        self.rules_tree.insert("", tk.END, values=(old, new))
        self.log(f"添加规则: '{old}' → '{new}'")
        
        # 清空输入
        self.old_word.set("")
        self.new_word.set("")
    
    def delete_rule(self):
        """删除选中的规则"""
        selected = self.rules_tree.selection()
        if not selected:
            return
        
        for item in selected:
            index = self.rules_tree.index(item)
            old, new = self.replace_rules[index]
            self.replace_rules.pop(index)
            self.rules_tree.delete(item)
            self.log(f"删除规则: '{old}' → '{new}'")
    
    def log(self, message):
        """添加日志"""
        self.log_text.configure(state=tk.NORMAL)
        self.log_text.insert(tk.END, message + "\n")
        self.log_text.see(tk.END)
        self.log_text.configure(state=tk.DISABLED)
    
    def start_replace(self):
        """开始替换"""
        folder = self.folder_path.get()
        if not folder:
            messagebox.showerror("错误", "请先选择文件夹")
            return
        
        if not self.replace_rules:
            messagebox.showerror("错误", "请添加至少一条替换规则")
            return
        
        if not os.path.isdir(folder):
            messagebox.showerror("错误", "选择的文件夹不存在")
            return
        
        # 创建输出文件夹
        if self.output_option.get() == "new_folder":
            output_folder = os.path.join(folder, "替换后文件")
            os.makedirs(output_folder, exist_ok=True)
        else:
            output_folder = folder
        
        self.log("=" * 50)
        self.log("开始处理...")
        
        # 支持的文件扩展名（仅新格式）
        supported_extensions = ['.docx', '.xlsx']
        # 旧格式提示
        old_extensions = ['.doc', '.xls']
        
        processed = 0
        errors = 0
        skipped_old = 0
        
        # 先检查是否有旧格式文件
        old_files = []
        for root, dirs, files in os.walk(folder):
            if "替换后文件" in root:
                continue
            for filename in files:
                ext = os.path.splitext(filename)[1].lower()
                if ext in old_extensions:
                    old_files.append(filename)
        
        if old_files:
            self.log(f"⚠️ 发现 {len(old_files)} 个旧格式文件，已跳过:")
            for f in old_files[:5]:  # 只显示前5个
                self.log(f"   {f}")
            if len(old_files) > 5:
                self.log(f"   ... 还有 {len(old_files) - 5} 个")
            self.log("提示: 请用Office打开后另存为 .docx/.xlsx 格式")
        
        # 处理支持的文件
        for root, dirs, files in os.walk(folder):
            if "替换后文件" in root:
                continue
            
            for filename in files:
                ext = os.path.splitext(filename)[1].lower()
                
                if ext in old_extensions:
                    skipped_old += 1
                    continue
                
                if ext not in supported_extensions:
                    continue
                
                filepath = os.path.join(root, filename)
                
                try:
                    if ext == '.docx':
                        self.replace_docx(filepath, output_folder)
                    elif ext == '.xlsx':
                        self.replace_xlsx(filepath, output_folder)
                    
                    processed += 1
                    self.log(f"✓ {filename}")
                except Exception as e:
                    errors += 1
                    self.log(f"✗ {filename}: {str(e)}")
        
        self.log("=" * 50)
        summary = f"处理完成！成功: {processed}, 失败: {errors}"
        if skipped_old > 0:
            summary += f", 跳过旧格式: {skipped_old}"
        self.log(summary)
        
        if skipped_old > 0:
            messagebox.showwarning("完成", f"{summary}\n\n⚠️ 有 {skipped_old} 个旧格式文件(.doc/.xls)被跳过\n请用Office另存为新格式后重试")
        else:
            messagebox.showinfo("完成", summary)
    
    def replace_docx(self, filepath, output_folder):
        """替换.docx文件中的关键词"""
        if not HAS_DOCX:
            raise Exception("缺少python-docx库")
        
        doc = Document(filepath)
        
        for old, new in self.replace_rules:
            # 替换段落中的文本
            for paragraph in doc.paragraphs:
                if old in paragraph.text:
                    # 需要保留格式，逐个run处理
                    for run in paragraph.runs:
                        if old in run.text:
                            run.text = run.text.replace(old, new)
            
            # 替换表格中的文本
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if old in cell.text:
                            for paragraph in cell.paragraphs:
                                for run in paragraph.runs:
                                    if old in run.text:
                                        run.text = run.text.replace(old, new)
        
        # 保存文件
        filename = os.path.basename(filepath)
        output_path = os.path.join(output_folder, filename)
        doc.save(output_path)
    
    def replace_xlsx(self, filepath, output_folder):
        """替换.xlsx文件中的关键词"""
        if not HAS_OPENPYXL:
            raise Exception("缺少openpyxl库")
        
        wb = load_workbook(filepath)
        
        for sheet in wb.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, str):
                        for old, new in self.replace_rules:
                            if old in cell.value:
                                cell.value = cell.value.replace(old, new)
        
        filename = os.path.basename(filepath)
        output_path = os.path.join(output_folder, filename)
        wb.save(output_path)


def main():
    root = tk.Tk()
    app = BatchReplacerApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
